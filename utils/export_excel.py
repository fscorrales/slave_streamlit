__all__ = [
    "build_comprobante_xlsx",
]

from datetime import date, datetime
from io import BytesIO
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from utils.transform_data import formato_moneda_ar

TITULO_COMPROBANTE: str = "Comprobante de Honorarios"
CANTIDAD_COLUMNAS: int = 3

# Campos del comprobante: (etiqueta mostrada, clave en ``main_data``)
CAMPOS_COMPROBANTE: list[tuple[str, str]] = [
    ("Ejercicio", "ejercicio"),
    ("Mes", "mes"),
    ("Fecha", "fecha"),
    ("N° Comprobante", "nro_comprobante"),
    ("Tipo", "tipo"),
    ("Importe Bruto", "importe_bruto"),
]

_FONT_TITULO = Font(bold=True, size=14)
_FONT_SECCION = Font(bold=True, size=11, color="FFFFFF")
_FILL_SECCION = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
_FONT_CABECERA = Font(bold=True)
_FILL_CABECERA = PatternFill(
    start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"
)
_THIN_SIDE = Side(style="thin", color="9CA3AF")
_BORDER_CELDA = Border(
    left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE
)


# --------------------------------------------------
def _texto(valor: object) -> str:
    """Convierte un valor escalar a ``str`` vacío si es ``None``/``NaN``."""
    if valor is None or pd.isna(valor):
        return ""
    return str(valor)


# --------------------------------------------------
def _moneda(valor: object) -> str:
    """Formatea un importe con el formato monetario AR de la vista."""
    if valor is None or pd.isna(valor):
        return formato_moneda_ar(0)
    return formato_moneda_ar(valor)  # type: ignore[arg-type]


# --------------------------------------------------
def _formato_fecha(valor: object) -> str:
    """Formatea fechas como ``DD/MM/YYYY`` (acepta ``str`` como respaldo)."""
    if isinstance(valor, (datetime, date)):
        return valor.strftime("%d/%m/%Y")
    return _texto(valor)


# --------------------------------------------------
def _valor_comprobante(campo: str, valor: object) -> str:
    """Normaliza un campo de ``main_data`` para mostrarlo en la hoja."""
    if campo == "fecha":
        return _formato_fecha(valor)
    if campo == "importe_bruto":
        return _moneda(valor)
    return _texto(valor)


# --------------------------------------------------
def _escribir_titulo(ws: Worksheet) -> int:
    """Escribe el título general del comprobante. Devuelve la próxima fila."""
    celda = ws.cell(row=1, column=1, value=TITULO_COMPROBANTE)
    celda.font = _FONT_TITULO
    celda.alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=CANTIDAD_COLUMNAS)
    ws.row_dimensions[1].height = 22
    return 3  # dejamos una fila en blanco debajo del título


# --------------------------------------------------
def _escribir_titulo_seccion(ws: Worksheet, fila: int, texto: str) -> int:
    """Escribe un encabezado de sección fusionado. Devuelve la próxima fila."""
    for columna in range(1, CANTIDAD_COLUMNAS + 1):
        ws.cell(row=fila, column=columna).fill = _FILL_SECCION
    celda = ws.cell(row=fila, column=1, value=texto)
    celda.font = _FONT_SECCION
    celda.alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells(
        start_row=fila, start_column=1, end_row=fila, end_column=CANTIDAD_COLUMNAS
    )
    ws.row_dimensions[fila].height = 20
    return fila + 1


# --------------------------------------------------
def _escribir_encabezado_columnas(
    ws: Worksheet, fila: int, cabeceras: list[str]
) -> int:
    """Escribe la fila de cabeceras de una tabla. Devuelve la próxima fila."""
    for columna, texto in enumerate(cabeceras, start=1):
        celda = ws.cell(row=fila, column=columna, value=texto)
        celda.font = _FONT_CABECERA
        celda.fill = _FILL_CABECERA
        celda.border = _BORDER_CELDA
        celda.alignment = Alignment(horizontal="center", vertical="center")
    return fila + 1


# --------------------------------------------------
def _escribir_fila_valores(
    ws: Worksheet,
    fila: int,
    valores: list[str],
    alinear_derecha: tuple[int, ...] = (),
    es_total: bool = False,
) -> int:
    """Escribe una fila de datos con bordes. Devuelve la próxima fila."""
    for columna, texto in enumerate(valores, start=1):
        celda = ws.cell(row=fila, column=columna, value=texto)
        celda.border = _BORDER_CELDA
        if es_total:
            celda.font = _FONT_CABECERA
        if columna in alinear_derecha:
            celda.alignment = Alignment(horizontal="right")
    return fila + 1


# --------------------------------------------------
def build_comprobante_xlsx(
    main_data: dict[str, Any],
    df_imp: pd.DataFrame,
    df_ret: pd.DataFrame,
) -> bytes:
    """
    Genera el archivo .xlsx de un comprobante de honorarios.

    La hoja única replica la disposición de la vista ``honorarios``:
    1. Datos del comprobante (``main_data``).
    2. Importes agrupados por actividad y partida (``df_imp``).
    3. Retenciones (``df_ret``).

    Args:
        main_data: Datos generales del comprobante (ejercicio, mes, fecha,
            nro_comprobante, tipo, importe_bruto).
        df_imp: DataFrame con columnas ``actividad``, ``partida`` e
            ``importe_bruto``.
        df_ret: DataFrame con columnas ``codigo`` e ``importe``. Puede estar
            vacío (se mostrará "Sin retenciones").

    Returns:
        Contenido binario del ``.xlsx`` listo para ``st.download_button``.
    """
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Comprobante"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 20

    fila = _escribir_titulo(ws)

    # 1. Datos del comprobante
    fila = _escribir_titulo_seccion(ws, fila, "Datos del comprobante")
    for etiqueta, campo in CAMPOS_COMPROBANTE:
        etiqueta_celda = ws.cell(row=fila, column=1, value=etiqueta)
        etiqueta_celda.font = _FONT_CABECERA
        etiqueta_celda.border = _BORDER_CELDA
        valor_celda = ws.cell(
            row=fila, column=2, value=_valor_comprobante(campo, main_data.get(campo))
        )
        valor_celda.border = _BORDER_CELDA
        fila += 1
    fila += 1  # fila en blanco entre secciones

    # 2. Importes por actividad y partida
    fila = _escribir_titulo_seccion(ws, fila, "Importes por actividad y partida")
    fila = _escribir_encabezado_columnas(ws, fila, ["Actividad", "Partida", "Importe"])
    for _, registro in df_imp.iterrows():
        fila = _escribir_fila_valores(
            ws,
            fila,
            [
                _texto(registro.get("actividad")),
                _texto(registro.get("partida")),
                _moneda(registro.get("importe_bruto")),
            ],
            alinear_derecha=(3,),
        )
    total_importes = df_imp["importe_bruto"].sum() if not df_imp.empty else 0
    fila = _escribir_fila_valores(
        ws,
        fila,
        ["Total", "", _moneda(total_importes)],
        alinear_derecha=(3,),
        es_total=True,
    )
    fila += 1  # fila en blanco entre secciones

    # 3. Retenciones
    fila = _escribir_titulo_seccion(ws, fila, "Retenciones")
    fila = _escribir_encabezado_columnas(ws, fila, ["Código", "Importe"])
    if df_ret.empty:
        fila = _escribir_fila_valores(ws, fila, ["Sin retenciones", ""])
    else:
        for _, registro in df_ret.iterrows():
            fila = _escribir_fila_valores(
                ws,
                fila,
                [_texto(registro.get("codigo")), _moneda(registro.get("importe"))],
                alinear_derecha=(2,),
            )
        total_retenciones = df_ret["importe"].sum()
        fila = _escribir_fila_valores(
            ws,
            fila,
            ["Total", _moneda(total_retenciones)],
            alinear_derecha=(2,),
            es_total=True,
        )

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
